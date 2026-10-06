# -*- coding: utf-8 -*-
import os
import sys
import tempfile
import unittest
import zipfile
from xml.dom import minidom

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from gcb.xlsx import Workbook, _col  # noqa: E402


class XlsxTest(unittest.TestCase):
    def build(self):
        wb = Workbook()
        ws = wb.sheet("摘要", widths=[20, 40])
        ws.row([("標題", "title")])
        ws.header(["項目", "值"])
        ws.row(["合格", (5, "pass")])
        ws.row(["含特殊字元", "<a & b>\x01\n第二行"])
        wb.sheet("摘要")  # 重複名稱自動改名
        wb.sheet("名稱:有/非法*字元?")
        return wb

    def test_col(self):
        self.assertEqual([_col(0), _col(25), _col(26), _col(701)], ["A", "Z", "AA", "ZZ"])

    def test_valid_xml_parts(self):
        path = os.path.join(tempfile.mkdtemp(), "t.xlsx")
        self.build().save(path)
        with zipfile.ZipFile(path) as z:
            for n in z.namelist():
                if n.endswith((".xml", ".rels")):
                    minidom.parseString(z.read(n))  # 解析失敗會丟例外

    def test_long_text_with_entities_stays_valid(self):
        wb = Workbook()
        wb.sheet("s").row(["a" + "<" * 20000])
        path = os.path.join(tempfile.mkdtemp(), "t.xlsx")
        wb.save(path)
        with zipfile.ZipFile(path) as z:
            minidom.parseString(z.read("xl/worksheets/sheet1.xml"))

    def test_openpyxl_reads(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("未安裝 openpyxl")
        path = os.path.join(tempfile.mkdtemp(), "t.xlsx")
        self.build().save(path)
        wb = openpyxl.load_workbook(path)
        self.assertEqual(wb.sheetnames[0], "摘要")
        self.assertEqual(len(wb.sheetnames), 3)
        ws = wb["摘要"]
        self.assertEqual(ws["B3"].value, 5)
        self.assertEqual(ws["B4"].value, "<a & b>\n第二行")
        self.assertEqual(ws.auto_filter.ref, "A2:B4")


if __name__ == "__main__":
    unittest.main()
